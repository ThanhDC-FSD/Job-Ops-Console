import argparse
import re
import sys
from pathlib import Path
from typing import Pattern

try:
    from docx import Document
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.shared import Pt, RGBColor
except ImportError:
    print("Missing dependency: python-docx")
    print("Install with: pip install python-docx")
    sys.exit(1)

try:
    from docx2pdf import convert as convert_docx2pdf

    DOCX2PDF_AVAILABLE = True
except ImportError:
    DOCX2PDF_AVAILABLE = False

from technical_skill_groups import TECHNICAL_SKILL_GROUPS, TECHNICAL_SKILL_GROUP_COLORS


TAG_PATTERN = re.compile(
    r"<(/?)(bold|italic|size|[a-zA-Z]+)(?::([0-9]+(?:\.[0-9]+)?))?>"
)
DEFAULT_COLOR_MAP = {
    "orange": RGBColor(230, 126, 34),
    "green": RGBColor(39, 174, 96),
    "blue": RGBColor(41, 128, 185),
}
COLOR_ENV_PREFIX = "COLOR_"
DEFAULT_FONT_SIZE = 11.0
DEFAULT_FONT_NAME = "Calibri"
FONT_NAME_ENV_KEY = "FONT_NAME"
PARAGRAPH_ALIGNMENT_ENV_KEY = "PARAGRAPH_ALIGNMENT"
DEFAULT_PARAGRAPH_ALIGNMENT = "justify"
PARAGRAPH_ALIGNMENT_MAP = {
    "left": WD_PARAGRAPH_ALIGNMENT.LEFT,
    "center": WD_PARAGRAPH_ALIGNMENT.CENTER,
    "right": WD_PARAGRAPH_ALIGNMENT.RIGHT,
    "justify": WD_PARAGRAPH_ALIGNMENT.JUSTIFY,
}
ALIGNMENT_TAG_MAP = {
    "left": WD_PARAGRAPH_ALIGNMENT.LEFT,
    "center": WD_PARAGRAPH_ALIGNMENT.CENTER,
    "right": WD_PARAGRAPH_ALIGNMENT.RIGHT,
    "justify": WD_PARAGRAPH_ALIGNMENT.JUSTIFY,
}

ENV_FILE = Path(".env")
INPUT_DIR = Path("input")
DEFAULT_INPUT_PATH = INPUT_DIR / "full_doc_stlye.txt"
FONT_SIZE_DEFAULTS = {
    "default": DEFAULT_FONT_SIZE,
    "section": 14.0,
    "role": 12.0,
    "date": 10.5,
    "list_item": 11.0,
    "summary_note": 10.0,
}
FONT_SIZE_ENV_KEYS = {
    "default": "DEFAULT_FONT_SIZE",
    "section": "SECTION_HEADER_FONT_SIZE",
    "role": "ROLE_TITLE_FONT_SIZE",
    "date": "ROLE_DATE_FONT_SIZE",
    "list_item": "LIST_ITEM_FONT_SIZE",
    "summary_note": "PROFESSIONAL_SUMMARY_FONT_SIZE",
}


def parse_env_file(env_path: Path):
    """Return simple key=value pairs from an .env file."""
    if not env_path.exists():
        return {}

    values = {}
    for line in env_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        values[key.strip()] = value.strip()
    return values


def load_font_sizes(env_values: dict[str, str]):
    """Resolve font sizes from .env overrides."""
    sizes = {}
    for key, env_key in FONT_SIZE_ENV_KEYS.items():
        raw_value = env_values.get(env_key)
        if raw_value:
            try:
                sizes[key] = float(raw_value)
                continue
            except ValueError:
                pass
        sizes[key] = FONT_SIZE_DEFAULTS[key]
    return sizes


def load_font_name(env_values: dict[str, str]) -> str:
    """Resolve font family from .env override."""
    font_name = env_values.get(FONT_NAME_ENV_KEY, "").strip()
    return font_name or DEFAULT_FONT_NAME


def load_paragraph_alignment(env_values: dict[str, str]):
    """Resolve paragraph alignment from .env override."""
    raw_alignment = env_values.get(
        PARAGRAPH_ALIGNMENT_ENV_KEY, DEFAULT_PARAGRAPH_ALIGNMENT
    ).strip().lower()
    return PARAGRAPH_ALIGNMENT_MAP.get(
        raw_alignment, PARAGRAPH_ALIGNMENT_MAP[DEFAULT_PARAGRAPH_ALIGNMENT]
    )


def load_color_map(env_values: dict[str, str]):
    """Build color map by combining defaults and .env overrides."""
    color_map: dict[str, RGBColor] = {}
    for key, raw_value in env_values.items():
        if not key.startswith(COLOR_ENV_PREFIX):
            continue
        color_name = key[len(COLOR_ENV_PREFIX) :].lower()
        tokens = [token.strip() for token in raw_value.split(",") if token.strip()]
        if len(tokens) != 3:
            continue
        try:
            rgb = [max(0, min(255, int(token))) for token in tokens]
        except ValueError:
            continue
        color_map[color_name] = RGBColor(*rgb)
    for name, color in DEFAULT_COLOR_MAP.items():
        color_map.setdefault(name, color)
    return color_map


def build_skill_color_map() -> dict[str, str]:
    """Map each technical skill term to the color name configured for its group."""
    skill_colors: dict[str, str] = {}
    for group_name, skills in TECHNICAL_SKILL_GROUPS.items():
        color_name = TECHNICAL_SKILL_GROUP_COLORS.get(group_name)
        if not color_name:
            continue
        for skill in skills:
            normalized = skill.strip().lower()
            if normalized:
                skill_colors[normalized] = color_name
    return skill_colors


def build_skill_pattern(skill_color_map: dict[str, str]) -> Pattern[str] | None:
    """Compile regex for matching all technical skills in plain text runs."""
    if not skill_color_map:
        return None
    escaped = sorted((re.escape(term) for term in skill_color_map), key=len, reverse=True)
    if not escaped:
        return None
    # Match complete skill tokens only, so "Java" does not match inside "JavaScript".
    return re.compile(
        r"(?<![A-Za-z0-9_])(?:"
        + "|".join(escaped)
        + r")(?![A-Za-z0-9_])",
        re.IGNORECASE,
    )


def split_text_by_skill(
    text: str,
    base_color_name: str | None,
    skill_pattern: Pattern[str] | None,
    skill_color_map: dict[str, str],
):
    """
    Split a segment into (substring, color_name) chunks.
    Skill tokens use group-based color, non-skill text keeps base color.
    """
    if not text or skill_pattern is None:
        return [(text, base_color_name)]

    chunks: list[tuple[str, str | None]] = []
    cursor = 0
    for match in skill_pattern.finditer(text):
        start, end = match.span()
        if start > cursor:
            chunks.append((text[cursor:start], base_color_name))

        token = match.group(0)
        token_color = skill_color_map.get(token.lower(), base_color_name)
        chunks.append((token, token_color))
        cursor = end

    if cursor < len(text):
        chunks.append((text[cursor:], base_color_name))

    return [(chunk_text, chunk_color) for chunk_text, chunk_color in chunks if chunk_text]


def determine_line_font_size(line: str, sizes: dict[str, float]) -> float:
    stripped = line.lstrip()
    if not stripped:
        return sizes["default"]
    if stripped.startswith("- "):
        return sizes["list_item"]
    if "<green>" in stripped and "<bold>" in stripped:
        return sizes["section"]
    if "|" in stripped:
        return sizes["role"]
    if stripped.startswith("<italic>") and stripped.endswith("</italic>"):
        return sizes["date"]
    return sizes["default"]


def parse_tagged_text(
    text: str, default_font_size: float, color_map: dict[str, RGBColor]
):
    """Return a list of (segment_text, is_bold, is_italic, color_name, font_size_pt)."""
    counters = {
        "bold": 0,
        "italic": 0,
    }
    color_stack = []
    size_stack = []
    segments = []
    cursor = 0

    for match in TAG_PATTERN.finditer(text):
        start, end = match.span()
        if start > cursor:
            segment = text[cursor:start]
            active_color = color_stack[-1] if color_stack else None
            segments.append(
                (
                    segment,
                    counters["bold"] > 0,
                    counters["italic"] > 0,
                    active_color,
                    size_stack[-1] if size_stack else default_font_size,
                )
            )

        is_closing = match.group(1) == "/"
        tag = match.group(2)
        size_value = match.group(3)
        if is_closing:
            if tag in counters:
                counters[tag] = max(0, counters[tag] - 1)
            elif tag in color_map:
                for i in range(len(color_stack) - 1, -1, -1):
                    if color_stack[i] == tag:
                        del color_stack[i]
                        break
            elif tag == "size" and size_stack:
                size_stack.pop()
        else:
            if tag in counters:
                counters[tag] += 1
            elif tag in color_map:
                color_stack.append(tag)
            elif tag == "size":
                if size_value is not None:
                    size_stack.append(float(size_value))
        cursor = end

    if cursor < len(text):
        segment = text[cursor:]
        active_color = color_stack[-1] if color_stack else None
        segments.append(
            (
                segment,
                counters["bold"] > 0,
                counters["italic"] > 0,
                active_color,
                size_stack[-1] if size_stack else default_font_size,
            )
        )

    return [s for s in segments if s[0]]


def add_styled_runs(
    paragraph,
    text: str,
    default_font_size: float,
    color_map: dict[str, RGBColor],
    font_name: str,
    skill_pattern: Pattern[str] | None,
    skill_color_map: dict[str, str],
    enable_skill_coloring: bool = True,
    force_italic: bool = False,
):
    for part, is_bold, is_italic, color_name, font_size in parse_tagged_text(
        text, default_font_size, color_map
    ):
        # Respect explicit color tags first; only auto-color plain (untagged) text.
        if enable_skill_coloring and color_name is None:
            chunks = split_text_by_skill(part, color_name, skill_pattern, skill_color_map)
        else:
            chunks = [(part, color_name)]
        for chunk_text, chunk_color in chunks:
            run = paragraph.add_run(chunk_text)
            run.bold = is_bold
            run.italic = is_italic or force_italic
            run.font.size = Pt(font_size)
            run.font.name = font_name
            # Auto-color "label:" tokens in bold (e.g., Languages:) when no explicit color tag is set.
            if chunk_color is None and is_bold and re.fullmatch(r"\s*[^:]+:\s*", chunk_text):
                chunk_color = "green"
            if chunk_color:
                color_value = color_map.get(chunk_color)
                if color_value:
                    run.font.color.rgb = color_value


def resolve_line_alignment(line: str, default_alignment):
    """Resolve paragraph alignment from inline tags like <center>...</center>."""
    lowered = line.lower()
    for tag_name, alignment in ALIGNMENT_TAG_MAP.items():
        if f"<{tag_name}>" in lowered:
            return alignment
    return default_alignment


def convert_text_to_docx(input_path: Path, output_path: Path):
    document = Document()

    env_values = parse_env_file(ENV_FILE)
    font_sizes = load_font_sizes(env_values)
    font_name = load_font_name(env_values)
    paragraph_alignment = load_paragraph_alignment(env_values)
    color_map = load_color_map(env_values)
    skill_color_map = build_skill_color_map()
    skill_pattern = build_skill_pattern(skill_color_map)
    normal_style = document.styles["Normal"]
    normal_style.font.name = font_name
    output_path.parent.mkdir(parents=True, exist_ok=True)
    # Ignore stray non-UTF8 bytes (common in downloaded/LLM text) to keep batch renders from failing.
    lines = input_path.read_text(encoding="utf-8", errors="ignore").splitlines()
    in_professional_summary = False
    for line in lines:
        raw_line = line.rstrip()
        if not raw_line:
            document.add_paragraph("")
            continue

        stripped = raw_line.lstrip()
        plain_text = re.sub(r"<[^>]+>", "", stripped).strip().lower()
        is_professional_summary_header = (
            "professional summary" in plain_text and not stripped.startswith("- ")
        )
        if is_professional_summary_header:
            in_professional_summary = True
        elif plain_text and not stripped.startswith("- "):
            in_professional_summary = False

        line_font_size = determine_line_font_size(raw_line, font_sizes)
        if in_professional_summary and stripped.startswith("- "):
            line_font_size = font_sizes.get("summary_note", line_font_size)

        line_alignment = resolve_line_alignment(raw_line, paragraph_alignment)
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.alignment = line_alignment
            add_styled_runs(
                paragraph,
                stripped[2:],
                line_font_size,
                color_map,
                font_name,
                skill_pattern,
                skill_color_map,
                enable_skill_coloring=True,
                force_italic=in_professional_summary,
            )
        else:
            paragraph = document.add_paragraph()
            paragraph.alignment = line_alignment
            is_role_title_line = "|" in stripped
            add_styled_runs(
                paragraph,
                raw_line,
                line_font_size,
                color_map,
                font_name,
                skill_pattern,
                skill_color_map,
                enable_skill_coloring=not is_role_title_line,
            )

    document.save(output_path)


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path):
    if not DOCX2PDF_AVAILABLE:
        raise RuntimeError(
            "Missing dependency: docx2pdf (install with: pip install docx2pdf)"
        )
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        convert_docx2pdf(str(docx_path), str(pdf_path))
    except Exception as exc:
        raise RuntimeError(f"Failed to export PDF: {exc}") from exc


def main():
    parser = argparse.ArgumentParser(
        description="Parse tagged CV text and export a styled .docx file."
    )
    parser.add_argument(
        "-i",
        "--input",
        default=str(DEFAULT_INPUT_PATH),
        help=f"Input tagged text file (default: {DEFAULT_INPUT_PATH})",
    )
    parser.add_argument(
        "-o",
        "--output",
        default="full_doc_stlye.docx",
        help="Output .docx file (default: full_doc_stlye.docx)",
    )
    parser.add_argument(
        "--pdf-output",
        default="",
        help="Optional PDF output path (default: same name as --output with .pdf)",
    )
    parser.add_argument(
        "--skip-pdf",
        action="store_true",
        help="Skip PDF export step.",
    )
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        print(f"Input file not found: {input_path}")
        sys.exit(1)

    convert_text_to_docx(input_path, output_path)
    print(f"Generated: {output_path}")

    if not args.skip_pdf:
        pdf_output_path = (
            Path(args.pdf_output) if args.pdf_output else output_path.with_suffix(".pdf")
        )
        try:
            convert_docx_to_pdf(output_path, pdf_output_path)
        except RuntimeError as exc:
            print(exc)
            sys.exit(1)
        print(f"Generated: {pdf_output_path}")


if __name__ == "__main__":
    main()
